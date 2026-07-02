"""Excel daily report generation."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from app.core.config import Settings
from app.core.logging import get_logger

logger = get_logger(__name__)

TRADE_COLUMNS = [
    "date", "time", "engine", "mode", "option_side", "strike", "expiry",
    "quantity", "lots", "entry_order_id", "exit_order_id", "entry_price",
    "exit_price", "points", "gross_pnl", "brokerage_estimate", "net_pnl",
    "exit_reason", "confidence", "market_mode",
]


class ExcelReport:
    def __init__(self, settings: Settings) -> None:
        self._reports_dir = settings.reports_dir

    def generate_daily(self, trades: list[dict], state: dict, filename: str | None = None) -> Path:
        self._reports_dir.mkdir(parents=True, exist_ok=True)
        date_str = datetime.now().strftime("%Y%m%d")
        path = self._reports_dir / (filename or f"daily_report_{date_str}.xlsx")

        wb = Workbook()
        self._write_trades_sheet(wb.active, "Combined Summary", trades)
        for engine in ("normal", "wick", "ultra"):
            ws = wb.create_sheet(engine.title() + " Engine")
            engine_trades = [t for t in trades if t.get("engine") == engine]
            self._write_trades_sheet(ws, engine, engine_trades)

        for log_name in ("Broker Log", "Execution Log", "Risk Log", "Error Log", "Audit Log", "Statistics"):
            wb.create_sheet(log_name)

        tmp = path.with_suffix(".xlsx.tmp")
        wb.save(tmp)
        tmp.replace(path)
        logger.info("Excel report saved: %s", path)
        return path

    def _write_trades_sheet(self, ws, title: str, trades: list[dict]) -> None:
        ws.title = title[:31]
        ws.append(TRADE_COLUMNS)
        for t in trades:
            entry = float(t.get("entry_price") or 0)
            exit_p = float(t.get("exit_price") or 0)
            qty = int(t.get("quantity") or 0)
            points = exit_p - entry if exit_p and entry else 0
            gross = float(t.get("pnl") or points * qty)
            brokerage = gross * 0.001
            ws.append([
                datetime.now().strftime("%Y-%m-%d"),
                t.get("closed_at", ""),
                t.get("engine", ""),
                "",
                t.get("option_side", ""),
                t.get("strike", ""),
                t.get("expiry", ""),
                qty,
                t.get("lots", ""),
                t.get("entry_order_id", ""),
                t.get("exit_order_id", ""),
                entry,
                exit_p,
                points,
                gross,
                brokerage,
                gross - brokerage,
                t.get("exit_reason", ""),
                "",
                "",
            ])
